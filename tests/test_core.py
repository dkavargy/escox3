"""Unit tests that do not need any model (taxonomy, validation, chunking, extraction, export)."""

import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from escox3.build_reference import nace_ancestors, nace_level  # noqa: E402
from escox3.export import to_csv  # noqa: E402
from escox3.llm.prompts import selection_schema  # noqa: E402
from escox3.retrieval import chunk_text  # noqa: E402
from escox3.taxonomy import ReferenceStore  # noqa: E402
from escox3.text_extraction import ExtractionError, extract_text  # noqa: E402
from escox3.validation import parse_json_object, validate_selection, verify_evidence  # noqa: E402


@pytest.fixture(scope="module")
def store():
    return ReferenceStore()


# ------------------------------------------------------------------ taxonomy / mapping layer
def test_reference_counts(store):
    c = store.manifest["counts"]
    assert c["green_skills"] == 570
    assert c["occupations"] == 3039
    assert c["isco_groups"] == 619
    assert c["nace_codes"] == 706


def test_every_occupation_has_valid_isco(store):
    for code in store.occupations.isco_unit:
        assert code and store.isco_group(code)["level"] == "unit"


def test_isco_derivation_matches_esco_code(store):
    for _, r in store.occupations.sample(200, random_state=0).iterrows():
        assert r.esco_code.startswith(r.isco_unit)


def test_label_resolution_never_invents(store):
    uri = store.green_skills.uri.iloc[0]
    assert store.resolve_label(uri)["label"] == store.green_skills.label.iloc[0]
    bad = store.resolve_label("http://data.europa.eu/esco/skill/does-not-exist")
    assert bad["resolved"] is False and bad["label"] is None
    assert store.resolve_label("http://data.europa.eu/esco/isco/C2149")["label"]


def test_green_status_undetermined_without_reference(store):
    if store.green_share is None and store.green_imported is None:
        uri = store.occupations.uri.iloc[0]
        assert store.occupation(uri)["green_status"]["value"] == "undetermined"


def test_nace_levels_and_multilingual(store):
    assert nace_level("D35.11") == "class" and nace_level("D35") == "division" and nace_level("D") == "section"
    assert nace_ancestors("C33.12") == ["C", "C33", "C33.1", "C33.12"]
    e = store.nace_entry("K62")
    assert e["label"] and e["level"] == "division"
    assert store.nace_title("K62", "de") != store.nace_title("K62", "en")
    h = store.nace_hierarchy("A01.13")
    assert [x["code"] for x in h] == ["A", "A01", "A01.1", "A01.13"]


# ------------------------------------------------------------------ validation
def test_parse_json_variants():
    assert parse_json_object('{"selected": []}')[0] == {"selected": []}
    assert parse_json_object('```json\n{"a": 1}\n```')[0] == {"a": 1}
    assert parse_json_object('Sure! {"a": {"b": "}"}} trailing')[0] == {"a": {"b": "}"}}
    assert parse_json_object("no json")[0] is None


def test_validation_rejects_hallucinations_duplicates_and_unsupported():
    text = "We need experience in renewable energy systems and energy efficiency."
    parsed = {"selected": [
        {"id": "S1", "support": "explicit", "evidence": "renewable energy systems"},
        {"id": "S1", "support": "explicit", "evidence": "renewable energy systems"},       # duplicate
        {"id": "S99", "support": "explicit", "evidence": "energy efficiency"},             # not a candidate
        {"id": "S2", "support": "sometimes", "evidence": "energy efficiency"},             # bad support value
        {"id": "S3", "support": "explicit", "evidence": "quantum blockchain synergy"},     # quote not in text
        {"id": "S4", "support": "implicit", "evidence": "energy efficiency"},
    ]}
    rep = validate_selection(parsed, ["S1", "S2", "S3", "S4"], text, max_items=10)
    assert [a["id"] for a in rep.accepted] == ["S1", "S4"]
    reasons = " | ".join(r["reason"] for r in rep.rejected)
    for key in ("duplicate", "id_not_in_candidates", "invalid support", "unsupported"):
        assert key in reasons
    assert rep.accepted[0]["evidence_span"] == [text.find("renewable"), text.find("renewable") + len("renewable energy systems")]


def test_validation_invalid_structure():
    rep = validate_selection({"foo": 1}, ["S1"], "x", 3)
    assert not rep.valid_json and rep.error


def test_schema_restricts_ids():
    s = selection_schema(["S1", "S2"], 3)
    assert s["properties"]["selected"]["items"]["properties"]["id"]["enum"] == ["S1", "S2"]


def test_evidence_fuzzy():
    ok, span, score = verify_evidence("environmental impact assessment", "Prepare environmental impact assessments for plants")
    assert ok and span


# ------------------------------------------------------------------ chunking
def test_chunker_word_boundaries():
    chunks = chunk_text("Information management and recycling, organic farming.")
    texts = [c.text for c in chunks]
    assert "Information management" in texts         # original ESCOX split "Inf|ormation" on "or"
    assert "recycling" in texts and "organic farming" in texts
    for c in chunks:
        assert "Information management and recycling, organic farming."[c.start:c.end] == c.text


# ------------------------------------------------------------------ text extraction
def test_txt_and_errors():
    t, info = extract_text("a.txt", "Solar engineer\r\nwith café".encode("utf-8"))
    assert t == "Solar engineer\nwith café"
    with pytest.raises(ExtractionError):
        extract_text("a.exe", b"xx")
    with pytest.raises(ExtractionError):
        extract_text("a.pdf", b"not a pdf")


def test_docx_and_pdf_roundtrip(tmp_path):
    import docx
    from reportlab.pdfgen import canvas

    d = docx.Document()
    d.add_paragraph("Wind turbine technician")
    tb = d.add_table(rows=1, cols=2)
    tb.cell(0, 0).text, tb.cell(0, 1).text = "Skill", "energy efficiency"
    buf = io.BytesIO()
    d.save(buf)
    t, info = extract_text("cv.docx", buf.getvalue())
    assert "Wind turbine technician" in t and "energy efficiency" in t

    pdf = tmp_path / "x.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "Photovoltaic system designer")
    c.save()
    t, info = extract_text("x.pdf", pdf.read_bytes())
    assert "Photovoltaic system designer" in t and info["pages"] == 1


# ------------------------------------------------------------------ export
def test_csv_export_contains_labels_and_ids():
    result = {"green_skills": [{"rank": 1, "label": "energy efficiency", "esco_uri": "http://x/s", "green_basis": "b",
                                "support": "explicit", "confidence": "high", "similarity": 0.7, "evidence": "e"}],
              "green_occupations": [], "sectors": [{"rank": 1, "label": "Hospital activities", "nace_code": "R86.1",
                                                    "level": "class", "hierarchy": [{"code": "R", "title": None}]}],
              "module_status": {"green_occupations": {"status": "none_found", "message": "No occupation confidently identified."}}}
    csv_text = to_csv(result)
    assert "energy efficiency" in csv_text and "http://x/s" in csv_text and "R86.1" in csv_text
    assert "No occupation confidently identified." in csv_text


# ------------------------------------------------------------------ optional green-occupation importers
def test_green_share_from_relations(tmp_path, store):
    from escox3.build_reference import build_green_occupations

    occ = store.occupations.uri.iloc[0]
    g1, g2 = store.green_skills.uri.iloc[0], store.green_skills.uri.iloc[1]
    rel = tmp_path / "occupationSkillRelations_en.csv"
    rel.write_text("occupationUri,relationType,skillType,skillUri\n"
                   f"{occ},essential,skill/competence,{g1}\n{occ},optional,knowledge,{g2}\n"
                   f"{occ},essential,skill/competence,http://data.europa.eu/esco/skill/x\n"
                   f"{occ},essential,skill/competence,http://data.europa.eu/esco/skill/y\n")
    df = build_green_occupations(store.green_skills, store.occupations, [], {}, rel)
    row = df.set_index("occupation_uri").loc[occ]
    assert row.n_green_all == 2 and row.n_skills_all == 4 and abs(row.green_share_all - 0.5) < 1e-9
    assert abs(row.green_share_essential - 1 / 3) < 1e-3


def test_missing_relations_reported_as_gap(tmp_path, store):
    from escox3.build_reference import build_green_occupations

    gaps = []
    assert build_green_occupations(store.green_skills, store.occupations, gaps, {}, tmp_path / "none.csv") is None
    assert gaps[0]["id"] == "green_occupation_reference_missing"
