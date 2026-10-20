"""Text extraction from uploaded documents (TXT, PDF, DOCX)."""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Tuple

log = logging.getLogger("escox3.text")

SUPPORTED = {".txt": "Plain text", ".text": "Plain text", ".md": "Markdown text", ".pdf": "PDF", ".docx": "Word (DOCX)"}


class ExtractionError(ValueError):
    """Raised with a user-friendly message when a document cannot be converted to text."""


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # de-hyphenate PDF line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> Tuple[str, dict]:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type '{ext or '?'}'. Please upload a TXT, PDF or DOCX file.")
    if not data:
        raise ExtractionError("The uploaded file is empty.")
    info = {"filename": Path(filename).name, "type": SUPPORTED[ext], "bytes": len(data)}

    if ext in (".txt", ".text", ".md"):
        for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                if enc == "utf-16" and "\x00" in text:
                    continue
                info["encoding"] = enc
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - latin-1 always decodes
            raise ExtractionError("Could not decode the text file.")

    elif ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ExtractionError("This file does not look like a valid PDF.")
        text, pages = "", 0
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = len(pdf.pages)
                text = "\n\n".join((p.extract_text() or "") for p in pdf.pages)
            info["extractor"] = "pdfplumber"
        except Exception as e:  # fall back to pypdf
            log.warning("pdfplumber failed (%s); trying pypdf", e)
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data))
                pages = len(reader.pages)
                text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
                info["extractor"] = "pypdf"
            except Exception as e2:
                raise ExtractionError(f"The PDF could not be read ({type(e2).__name__}). It may be encrypted or damaged.")
        info["pages"] = pages
        if not text.strip():
            raise ExtractionError("No selectable text was found in this PDF (it may be a scanned image). "
                                  "Please run OCR first or paste the text.")

    else:  # .docx
        try:
            import docx

            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise ExtractionError(f"The DOCX file could not be opened ({type(e).__name__}). Old .doc files are not supported.")
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and t not in cells:
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        info["paragraphs"] = len(d.paragraphs)
        info["tables"] = len(d.tables)

    text = _clean(text)
    if not text:
        raise ExtractionError("The document does not contain any text.")
    info["chars"] = len(text)
    return text, info
